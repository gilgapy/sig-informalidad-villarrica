"""Ajusta el informe generado por pandoc al formato FCE-UNVES (APA 7.ª ed.).

- Ancho de tabla al 100 % del área de texto con ajuste automático de columnas.
- Letra de 10 pt en tablas; en tablas de 8 o más columnas, 9 pt y márgenes de celda reducidos
  (8 pt desde 9 columnas, para que los encabezados no partan palabras).
- Filas de tabla que no se parten entre páginas; cada tabla se mantiene en una sola página
  y su fila de encabezado se repite si aun así ocupa más de una.
- Cada figura se mantiene en la misma página que su nota.
- Los párrafos que solo contienen un salto de página se reemplazan por "salto de página
  antes" en el párrafo siguiente (evita líneas vacías al inicio de página y páginas en blanco).

Uso: python3 postprocesar_docx.py <archivo.docx>
"""

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FULL_WIDTH_PCT = "5000"  # 100 % expresado en cincuentavos de punto porcentual
WIDE_TABLE_COLUMNS = 8
EXTRA_WIDE_TABLE_COLUMNS = 9
WIDE_TABLE_CELL_MARGIN_TWIPS = "45"  # ~0,03 pulgadas a cada lado
FONT_SIZE_DEFAULT = Pt(10)
FONT_SIZE_WIDE = Pt(9)
FONT_SIZE_EXTRA_WIDE = Pt(8)


# Orden de los hijos de w:tblPr según el esquema OOXML (CT_TblPr)
TBL_PR_ORDER = [
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
    "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
    "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
]


def replace_tbl_pr_child(tbl_pr, element) -> None:
    """Reemplaza el hijo de tblPr con la misma etiqueta respetando el orden del esquema."""
    for existing in tbl_pr.findall(element.tag):
        tbl_pr.remove(existing)
    tags = [qn(tag) for tag in TBL_PR_ORDER]
    position = tags.index(element.tag)
    following = [child for child in tbl_pr if child.tag in tags[position + 1:]]
    if following:
        following[0].addprevious(element)
    else:
        tbl_pr.append(element)


def set_full_width(table) -> None:
    tbl_pr = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), FULL_WIDTH_PCT)
    width.set(qn("w:type"), "pct")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    replace_tbl_pr_child(tbl_pr, width)
    replace_tbl_pr_child(tbl_pr, layout)


def set_narrow_cell_margins(table) -> None:
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), WIDE_TABLE_CELL_MARGIN_TWIPS)
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    replace_tbl_pr_child(table._tbl.tblPr, margins)


def set_font_size(table, size) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = size


def set_row_flag(row, tag: str) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn(tag)) is None:
        tr_pr.append(OxmlElement(tag))


def keep_table_on_one_page(table) -> None:
    rows = list(table.rows)
    for index, row in enumerate(rows):
        set_row_flag(row, "w:cantSplit")
        if index == 0:
            set_row_flag(row, "w:tblHeader")
        if index < len(rows) - 1:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def keep_figures_with_notes(document) -> int:
    figures = [p for p in document.paragraphs if p._p.xpath(".//w:drawing")]
    for paragraph in figures:
        paragraph.paragraph_format.keep_with_next = True
    return len(figures)


def is_page_break_only(paragraph_element) -> bool:
    runs = paragraph_element.findall(qn("w:r"))
    if len(runs) != 1 or paragraph_element.find(qn("w:pPr")) is not None:
        return False
    children = [child for child in runs[0] if child.tag != qn("w:rPr")]
    return (
        len(children) == 1
        and children[0].tag == qn("w:br")
        and children[0].get(qn("w:type")) == "page"
    )


def next_paragraph(element):
    """Párrafo siguiente, saltando solo marcadores; None si lo que sigue es una tabla u otro bloque."""
    skippable = {qn("w:bookmarkStart"), qn("w:bookmarkEnd")}
    sibling = element.getnext()
    while sibling is not None and sibling.tag in skippable:
        sibling = sibling.getnext()
    if sibling is None or sibling.tag != qn("w:p"):
        return None
    return sibling


def replace_page_break_paragraphs(document) -> int:
    body = document.element.body
    replaced = 0
    # De atrás hacia adelante: dos saltos seguidos no dejan un salto suelto en el párrafo siguiente
    for paragraph_element in reversed(list(body.iterchildren(qn("w:p")))):
        if not is_page_break_only(paragraph_element):
            continue
        following = next_paragraph(paragraph_element)
        if following is None:
            continue
        p_pr = following.get_or_add_pPr()
        if p_pr.find(qn("w:pageBreakBefore")) is None:
            # pageBreakBefore debe ir después de pStyle, keepNext y keepLines (orden del esquema)
            anchor = None
            for tag in ("w:pStyle", "w:keepNext", "w:keepLines"):
                found = p_pr.find(qn(tag))
                if found is not None:
                    anchor = found
            page_break = OxmlElement("w:pageBreakBefore")
            if anchor is None:
                p_pr.insert(0, page_break)
            else:
                anchor.addnext(page_break)
        body.remove(paragraph_element)
        replaced += 1
    return replaced


def font_size_for(columns: int):
    if columns >= EXTRA_WIDE_TABLE_COLUMNS:
        return FONT_SIZE_EXTRA_WIDE
    if columns >= WIDE_TABLE_COLUMNS:
        return FONT_SIZE_WIDE
    return FONT_SIZE_DEFAULT


def main(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {path}")
    document = Document(str(path))
    for table in document.tables:
        set_full_width(table)
        columns = len(table.columns)
        if columns >= WIDE_TABLE_COLUMNS:
            set_narrow_cell_margins(table)
        set_font_size(table, font_size_for(columns))
        keep_table_on_one_page(table)
    figures = keep_figures_with_notes(document)
    page_breaks = replace_page_break_paragraphs(document)
    document.save(str(path))
    print(f"Tablas ajustadas: {len(document.tables)}; figuras unidas a su nota: {figures}; "
          f"saltos de página convertidos: {page_breaks} en {path.name}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Uso: python3 postprocesar_docx.py <archivo.docx>")
    main(Path(sys.argv[1]))
